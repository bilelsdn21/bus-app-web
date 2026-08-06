"""Excel export — Calendrier (colored day grid + excursion detail + summary) and the
externally-rented-vehicles list. Uses openpyxl (already a dependency for reading
final.xlsm in seed.py), writing straight to an in-memory buffer — no temp files."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from .. import models
from . import calc
from .calendar import build_calendar
from .lookups import region_periods_by_bus, region_for

HEADER_FILL = PatternFill("solid", fgColor="1A3A5C")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REGION_FILL = PatternFill("solid", fgColor="0F6EAB")
REGION_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="1A3A5C")
TOTAL_FONT = Font(color="FFFFFF", bold=True)
OUT_FILL = PatternFill("solid", fgColor="E2E8F0")   # hors contrat
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Same palette as the app's day-color legend (colors.js DOT), so the export reads
# identically to the on-screen calendar.
DAY_FILL = {
    "green":  PatternFill("solid", fgColor="10B981"),
    "yellow": PatternFill("solid", fgColor="FBBF24"),
    "orange": PatternFill("solid", fgColor="F97316"),
    "purple": PatternFill("solid", fgColor="8B5CF6"),
    "red":    PatternFill("solid", fgColor="F43F5E"),
}
DAY_LABEL = {"green": "Résa", "yellow": "Matin", "orange": "Soir", "purple": "Nuit", "red": "Indispo"}


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _style_row(ws, row, fill, font):
    for c in ws[row]:
        c.fill = fill
        c.font = font


def _calendar_sheet(wb, data):
    ws = wb.active
    ws.title = "Calendrier"
    days_in_month = data["days_in_month"]
    headers = ["Véhicule", "Région"] + [str(d) for d in range(1, days_in_month + 1)] + ["Net Mois (TND)", "%", "Distance (km)"]
    ws.append(headers)
    _style_row(ws, 1, HEADER_FILL, HEADER_FONT)
    for c in ws[1]:
        c.alignment = Alignment(horizontal="center")

    last_region = None
    for bus in data["buses"]:
        if bus["region"] != last_region:
            last_region = bus["region"]
            ws.append([bus["region"] or "—"] + [""] * (len(headers) - 1))
            _style_row(ws, ws.max_row, REGION_FILL, REGION_FONT)
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(headers))

        covered = set(bus.get("covered_days") or [])
        row = [bus["name"], bus["region"] or "—"]
        for d in range(1, days_in_month + 1):
            color = bus["days"].get(d)
            if color:
                row.append(DAY_LABEL.get(color, color))
            elif covered and d not in covered:
                row.append("Hors ctr")
            else:
                row.append("")
        row += [bus.get("net"), bus.get("pct"), bus.get("distance") or 0]
        ws.append(row)
        r = ws.max_row
        for i, d in enumerate(range(1, days_in_month + 1), start=3):
            color = bus["days"].get(d)
            cell = ws.cell(row=r, column=i)
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER
            if color and color in DAY_FILL:
                cell.fill = DAY_FILL[color]
            elif covered and d not in covered:
                cell.fill = OUT_FILL
        ws.cell(row=r, column=len(headers) - 2).number_format = '#,##0" TND"'
        ws.cell(row=r, column=len(headers)).number_format = '#,##0" km"'

    _autosize(ws, [28, 10] + [6] * days_in_month + [13, 8, 13])
    ws.freeze_panes = "C2"


def _excursions_sheet(wb, db: Session, year: int, month: int):
    ws = wb.create_sheet("Excursions")
    start, end = calc.month_bounds(year, month)
    books = (db.query(models.Booking)
             .join(models.Bus, models.Booking.bus_id == models.Bus.id)
             .options(joinedload(models.Booking.bus))
             .filter(models.Booking.date <= end,
                     func.coalesce(models.Booking.end_date, models.Booking.date) >= start)
             .order_by(models.Booking.date, models.Bus.sort_order).all())

    periods_by_bus = region_periods_by_bus(db)
    headers = ["Date début", "Date fin", "Véhicule", "Région", "Type", "Destination", "Client", "Pax",
               "Heure début", "Heure fin", "Total (TND)"]
    ws.append(headers)
    _style_row(ws, 1, HEADER_FILL, HEADER_FONT)

    for b in books:
        bus = b.bus
        region = region_for(bus, start, end, periods_by_bus) if bus else ""
        ws.append([
            b.date.isoformat(), (b.end_date or b.date).isoformat(),
            bus.name if bus else f"#{b.bus_id}", region or "—",
            "Réservation" if b.type == "Booking" else "Indisponible",
            b.destination or "", b.client or "", b.pax or 0,
            b.heure_debut or "", b.heure_fin or "",
            round(b.total or 0, 2) if b.type == "Booking" else None,
        ])
        ws.cell(row=ws.max_row, column=11).number_format = '#,##0.00" TND"'

    _autosize(ws, [12, 12, 28, 10, 14, 24, 20, 6, 11, 11, 13])
    ws.freeze_panes = "A2"


def _summary_sheet(wb, data):
    ws = wb.create_sheet("Résumé")
    ws.append([f"Résumé — {data['label']}"])
    ws["A1"].font = Font(bold=True, size=14, color="1A3A5C")
    ws.append([])
    ws.append(["Région", "Loyers (TND)", "Net (TND)", "%"])
    _style_row(ws, ws.max_row, HEADER_FILL, HEADER_FONT)
    for r in data["summary"]["regions"]:
        ws.append([r["region"], round(r["loyer"], 2), round(r["net"], 2), r["pct"]])
    t = data["summary"]["total"]
    ws.append(["TOTAL", round(t["loyer"], 2), round(t["net"], 2), t["pct"]])
    _style_row(ws, ws.max_row, TOTAL_FILL, TOTAL_FONT)

    ws.append([])
    rentals = data["summary"]["rentals"]
    ws.append(["🚐 Véhicules loués en renfort (info — non déduit du Net ci-dessus)"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="B45309")
    if rentals["items"]:
        ws.append(["Véhicule", "Région", "Coût ce mois (TND)"])
        _style_row(ws, ws.max_row, PatternFill("solid", fgColor="FEF3C7"), Font(bold=True, color="92400E"))
        for it in rentals["items"]:
            ws.append([it["name"], it["region"] or "—", round(it["cost"], 2)])
        ws.append(["Total loués", "", round(rentals["total"], 2)])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    else:
        ws.append(["Aucun véhicule loué ce mois-ci."])

    _autosize(ws, [32, 16, 16, 10])


def build_calendar_workbook(db: Session, year: int, month: int) -> BytesIO:
    data = build_calendar(db, year, month)
    wb = Workbook()
    _calendar_sheet(wb, data)
    _excursions_sheet(wb, db, year, month)
    _summary_sheet(wb, data)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_rentals_workbook(db: Session) -> BytesIO:
    rows = db.query(models.RentedVehicle).order_by(models.RentedVehicle.start_date.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Véhicules loués"
    ws.append(["Véhicules loués — export complet"])
    ws["A1"].font = Font(bold=True, size=14, color="1A3A5C")
    ws.append([])
    ws.append(["Nom", "Région", "Du", "Au", "Coût (TND)", "Notes"])
    _style_row(ws, ws.max_row, HEADER_FILL, HEADER_FONT)

    total = 0.0
    for r in rows:
        ws.append([r.name, r.region or "—", r.start_date.isoformat(), r.end_date.isoformat(),
                   round(r.cost or 0, 2), r.notes or ""])
        ws.cell(row=ws.max_row, column=5).number_format = '#,##0.00" TND"'
        total += r.cost or 0
    ws.append(["TOTAL", "", "", "", round(total, 2), ""])
    _style_row(ws, ws.max_row, TOTAL_FILL, TOTAL_FONT)
    ws.cell(row=ws.max_row, column=5).number_format = '#,##0.00" TND"'

    _autosize(ws, [30, 12, 14, 14, 14, 32])
    ws.freeze_panes = "A4"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
